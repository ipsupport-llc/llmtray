#!/usr/bin/env python3
"""Completeness of `extract` on .xlsx / .pptx against openpyxl / python-pptx.
xlsx: share of non-empty cells (cached values) found as a cell of the same sheet.
pptx: share of the reference's words (slide shapes, tables, groups, notes) present.
usage: PYTHONPATH=<libs> ooxml_vs_ref.py FILE..."""
import datetime
import json
import os
import re
import subprocess
import sys

EXE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".build", "release", "extract")
WORD = re.compile(r"\w{2,}", re.UNICODE)


def run(f):
    p = subprocess.run([EXE, f], capture_output=True, timeout=60)
    return [json.loads(l) for l in p.stdout.decode().splitlines() if l.strip()]


def xlsx(f):
    import openpyxl
    wb = openpyxl.load_workbook(f, data_only=True, read_only=False)
    pages = {p.get("name"): p for p in run(f)}
    ref = hit = 0
    miss = []
    for ws in wb.worksheets:
        pg = pages.get(ws.title)
        text = pg["text"] if pg else ""
        cells = set(c.strip() for line in text.split("\n") for c in line.split(" | "))
        nums = set()
        for c in cells:
            try:
                nums.add(round(float(c), 6))
            except ValueError:
                pass
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                ref += 1
                if isinstance(v, bool):
                    ok = ("TRUE" if v else "FALSE") in cells
                elif isinstance(v, (int, float)):
                    ok = round(float(v), 6) in nums
                elif isinstance(v, datetime.datetime):
                    ok = any(k.startswith(v.strftime("%Y-%m-%d")) for k in cells) or v.strftime("%H:%M") in text
                elif isinstance(v, (datetime.date, datetime.time)):
                    ok = v.isoformat()[:10] in text or str(v)[:5] in text
                else:
                    s = str(v).replace("\n", " ").strip()
                    ok = s in cells or s in text
                hit += ok
                if not ok and len(miss) < 3:
                    miss.append((ws.title, c.coordinate, repr(v)[:40]))
    return ref, hit, miss


def pptx(f):
    from pptx import Presentation
    prs = Presentation(f)
    words = set()

    def shape_words(sh):
        if sh.shape_type == 6 and hasattr(sh, "shapes"):  # group
            for s in sh.shapes:
                shape_words(s)
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            words.update(w.lower() for w in WORD.findall(sh.text_frame.text))
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    words.update(w.lower() for w in WORD.findall(c.text))

    for s in prs.slides:
        for sh in s.shapes:
            shape_words(sh)
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            words.update(w.lower() for w in WORD.findall(s.notes_slide.notes_text_frame.text))
    got = set(w.lower() for p in run(f) for w in WORD.findall(p["text"]))
    missing = sorted(words - got)
    return len(words), len(words & got), missing[:6]


tr = th = 0
for f in sys.argv[1:]:
    try:
        ref, hit, miss = (xlsx if f.endswith(".xlsx") else pptx)(f)
    except Exception as e:  # noqa: BLE001
        print(f"{os.path.basename(f)[:36]:36s} reference failed: {e}")
        continue
    tr += ref
    th += hit
    print(f"{os.path.basename(f)[:36]:36s} ref={ref:6d} found={hit:6d} {100.0 * hit / max(ref, 1):6.2f}%  {miss if miss else ''}")
print(f"TOTAL {th}/{tr} = {100.0 * th / max(tr, 1):.2f}%")
